from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app import attachments
from app.attachments import AttachmentError, parse_attachment, parse_attachments


def pdf_bytes(text: str | None = None) -> bytes:
    output = BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    if text:
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
            }
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(output)
    return output.getvalue()


def test_code_extension_is_text_even_with_generic_mime():
    result = parse_attachment(
        "<template>你好</template>".encode(), "Component.vue", "application/octet-stream"
    )
    assert result["extracted_text"] == "<template>你好</template>"
    assert result["status"] == "extracted"


def test_text_limit_is_consistent_with_saved_preview():
    result = parse_attachment(("字" * 8001).encode(), "example.txt")
    assert len(result["extracted_text"]) == 8000
    assert result["truncated"] is True
    assert "8000" in result["note"]


def test_unicode_bom_is_decoded():
    assert parse_attachment("你好".encode("utf-16"), "example.txt")["extracted_text"] == "你好"


def test_docx_includes_table_text_in_reading_order():
    document = Document()
    document.add_paragraph("before")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "inside table"
    document.add_paragraph("after")
    output = BytesIO()
    document.save(output)
    result = parse_attachment(output.getvalue(), "example.docx")
    assert result["extracted_text"] == "before\ninside table\nafter"


def test_pdf_text_and_metadata_only_pdf():
    assert parse_attachment(pdf_bytes("Hello PDF"), "text.pdf")["extracted_text"] == "Hello PDF"
    result = parse_attachment(pdf_bytes(), "scan.pdf")
    assert result["status"] == "metadata_only"
    assert result["extracted_text"] == ""
    assert "扫描件" in result["note"]


def test_doc_has_explicit_metadata_only_note():
    result = parse_attachment(b"legacy word bytes", "old.doc")
    assert result["status"] == "metadata_only"
    assert "DOCX" in result["note"]


@pytest.mark.parametrize(
    "name,data",
    [
        ("bad.pdf", b"broken pdf"),
        ("bad.docx", b"broken zip"),
        ("bad.txt", b"\xff\x80"),
        ("binary.py", b"binary\x00data"),
        ("image.png", b"png"),
    ],
)
def test_bad_files_raise_instead_of_becoming_successful_empty_text(name, data):
    with pytest.raises(AttachmentError):
        parse_attachment(data, name)


def test_file_and_batch_limits_are_explicit(monkeypatch):
    monkeypatch.setattr(attachments, "MAX_FILE_BYTES", 4)
    monkeypatch.setattr(attachments, "MAX_TOTAL_BYTES", 6)
    with pytest.raises(AttachmentError) as error:
        parse_attachment(b"12345", "large.txt")
    assert error.value.code == "ATTACHMENT_TOO_LARGE"
    with pytest.raises(AttachmentError) as error:
        parse_attachments([("a.txt", None, b"1234"), ("b.txt", None, b"123")])
    assert error.value.code == "ATTACHMENTS_TOO_LARGE"
    with pytest.raises(AttachmentError) as error:
        parse_attachments([("a.txt", None, b"")] * 11)
    assert error.value.code == "TOO_MANY_ATTACHMENTS"
